
import s5_church

from openpyxl import Workbook

from openpyxl.chart import (
    PieChart,
    ProjectedPieChart,
    Reference
)
from openpyxl.chart.series import DataPoint

# Numbers - Grab data from the master church program
paidStaff = s5_church.paidStaff()            #
volunteerStaff = s5_church.volunteerStaff()    # This should be grabbing from s5_church length variable

data = [
    ['Category', 'Number'],
    ['Paid', paidStaff],
    ['Volunteer', volunteerStaff],
]

wb = Workbook()
ws = wb.active
ws = wb.create_sheet(title="Ministry Paid")


for row in data:
    ws.append(row)

pie = PieChart()
labels = Reference(ws, min_col=1, min_row=2, max_row=5)
data = Reference(ws, min_col=2, min_row=1, max_row=5)
pie.add_data(data, titles_from_data=True)
pie.set_categories(labels)
pie.title = "Volunteers vs. Paid"

# Cut the first slice out of the pie
slice = DataPoint(idx=0, explosion=20)
pie.series[0].data_points = [slice]

ws.add_chart(pie, "D1")




wb.save("pie.xlsx")
